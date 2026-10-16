#!/usr/bin/env python3
"""verify.py — gate an .xlsx before handoff.

    python3 verify.py out.xlsx            # static + LibreOffice recalc (if installed)
    python3 verify.py out.xlsx --no-calc  # static only
    python3 verify.py out.xlsx --json     # machine-readable

Exit 0 = clean, 1 = FAILs present. WARNs never fail the gate.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
NUM_RE = re.compile(r"^\(?[-+]?[₹$€£]?\s*\d{1,3}(,\d{2,3})*(\.\d+)?\s*%?\)?$|^\(?[-+]?\d+\.?\d*\)?%?$")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?$|^\d{2}[-/]\d{2}[-/]\d{4}$")
TOTAL_RE = re.compile(r"^\s*(grand\s+)?(total|totals|sum|subtotal|sub-total|overall|all\b.*)\s*[:\-]?\s*$", re.I)
DATE_TOKENS = ("y", "d", "h", "s")
SOFFICE_CANDIDATES = (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice",
    "/usr/local/bin/soffice",
)


def find_soffice() -> str | None:
    for p in SOFFICE_CANDIDATES:
        if os.path.exists(p):
            return p
    return shutil.which("soffice") or shutil.which("libreoffice")


def _is_id_like(s: str) -> bool:
    return s.isdigit() and len(s) > 1 and s[0] == "0"


def has_red(number_format: str | None) -> bool:
    """True if the format paints negatives red. LibreOffice round-trips [Red] as [RED], and
    some locales emit the indexed form [Color3] — both are the same instruction."""
    nf = (number_format or "").lower()
    return "[red]" in nf or "[color3]" in nf


def dead_totals(ws) -> list[str]:
    """Numbers that sit in a row or column labelled TOTAL/SUM/SUBTOTAL but are hardcoded."""
    hits, label_rows, label_cols = [], set(), set()
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and TOTAL_RE.match(c.value):
                if c.column <= 2:
                    label_rows.add(c.row)
                if c.row <= 3:
                    label_cols.add(c.column)
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, (int, float)) or isinstance(c.value, bool):
                continue
            if c.row in label_rows or c.column in label_cols:
                hits.append(c.coordinate)
    return hits[:20]


def static_check(path: str) -> dict:
    wb = load_workbook(path, data_only=False)
    out = {"fail": [], "warn": [], "info": {}, "sheets": {}}

    if "README" not in wb.sheetnames:
        out["fail"].append({"code": "no_readme", "msg": "No README sheet. Rule 6."})
    elif wb.sheetnames[0] != "README":
        out["warn"].append({"code": "readme_not_first", "msg": "README exists but is not sheet 1."})

    calc = bool(getattr(wb.calculation, "fullCalcOnLoad", False))
    out["info"]["fullCalcOnLoad"] = calc
    if not calc:
        out["fail"].append({"code": "no_fullcalc",
                            "msg": "wb.calculation.fullCalcOnLoad is not True — formulas may open blank."})

    for nm, dn in wb.defined_names.items():
        ref = str(getattr(dn, "value", "") or getattr(dn, "attr_text", ""))
        sheet_ref = ref.split("!")[0].strip("'") if "!" in ref else None
        if sheet_ref and sheet_ref not in wb.sheetnames:
            out["fail"].append({"code": "dead_name", "msg": f"Named range '{nm}' -> missing sheet '{sheet_ref}'"})
    out["info"]["named_ranges"] = len(wb.defined_names)

    for ws in wb.worksheets:
        s = {"formulas": 0, "constants": 0, "text_numbers": [], "text_dates": [],
             "general_format": [], "unformatted_dates": [], "negatives_not_red": [],
             "merged": len(ws.merged_cells.ranges), "rows": ws.max_row, "cols": ws.max_column}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith("="):
                    s["formulas"] += 1
                    if any(e in v for e in ERRORS):
                        out["fail"].append({"code": "error_in_formula",
                                            "msg": f"{ws.title}!{c.coordinate} formula contains an error literal: {v[:80]}"})
                    continue
                s["constants"] += 1
                if isinstance(v, str):
                    t = v.strip()
                    if any(e in t for e in ERRORS):
                        out["fail"].append({"code": "error_value",
                                            "msg": f"{ws.title}!{c.coordinate} holds error text {t[:40]}"})
                    elif c.row > 1 and NUM_RE.match(t) and not _is_id_like(t):
                        s["text_numbers"].append(c.coordinate)
                    elif c.row > 1 and DATE_RE.match(t):
                        s["text_dates"].append(c.coordinate)
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    nf = c.number_format or "General"
                    if nf in ("General", None):
                        s["general_format"].append(c.coordinate)
                    if v < 0 and not has_red(nf):
                        s["negatives_not_red"].append(c.coordinate)
                elif isinstance(v, (_dt.datetime, _dt.date)):
                    nf = (c.number_format or "").lower()
                    if not any(tok in nf for tok in DATE_TOKENS):
                        s["unformatted_dates"].append(c.coordinate)

        if s["text_numbers"]:
            out["fail"].append({"code": "numbers_as_text",
                                "msg": f"{ws.title}: {len(s['text_numbers'])} numeric values stored as text "
                                       f"(e.g. {', '.join(s['text_numbers'][:5])}). SUM will ignore them. Rule 3."})
        if s["text_dates"]:
            out["fail"].append({"code": "dates_as_text",
                                "msg": f"{ws.title}: {len(s['text_dates'])} dates stored as text "
                                       f"(e.g. {', '.join(s['text_dates'][:5])}). Rule 3."})
        if s["negatives_not_red"]:
            out["fail"].append({"code": "negative_not_red",
                                "msg": f"{ws.title}: {len(s['negatives_not_red'])} negative values whose "
                                       f"number_format has no [Red] section "
                                       f"(e.g. {', '.join(s['negatives_not_red'][:5])}). Rule 10a."})
        if s["unformatted_dates"]:
            out["warn"].append({"code": "date_no_format",
                                "msg": f"{ws.title}: {len(s['unformatted_dates'])} date cells with no date number_format."})
        if s["general_format"]:
            out["warn"].append({"code": "general_format",
                                "msg": f"{ws.title}: {len(s['general_format'])} numeric cells left on General "
                                       f"(no thousands separator / decimal discipline). Rule 4/10."})
        if s["merged"]:
            out["warn"].append({"code": "merged_cells",
                                "msg": f"{ws.title}: {s['merged']} merged ranges — breaks filtering and pivots."})
        for coord in dead_totals(ws):
            out["fail"].append({"code": "dead_total",
                                "msg": f"{ws.title}!{coord} sits in a row/column labelled a total but holds a "
                                       f"hardcoded number, not a formula. Rule 1."})

        if ws.title not in ("README", "Checks") and s["formulas"] == 0 and s["constants"] > 20:
            out["warn"].append({"code": "no_formulas",
                                "msg": f"{ws.title}: {s['constants']} values and ZERO formulas. "
                                       f"If any of these are totals or derived, that violates Rule 1."})
        s["text_numbers"] = s["text_numbers"][:20]
        s["text_dates"] = s["text_dates"][:20]
        s["general_format"] = s["general_format"][:20]
        s["unformatted_dates"] = s["unformatted_dates"][:20]
        s["negatives_not_red"] = s["negatives_not_red"][:20]
        out["sheets"][ws.title] = s

    wb.close()
    return out


def recalc_check(path: str) -> dict:
    soffice = find_soffice()
    if not soffice:
        return {"status": "skipped", "reason": "LibreOffice not found — static checks only."}
    with tempfile.TemporaryDirectory() as td:
        profile = os.path.join(td, "profile")
        cmd = [soffice, "--headless", "--norestore", "--invisible",
               f"-env:UserInstallation=file://{profile}",
               "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", td, os.path.abspath(path)]
        try:
            p = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        except subprocess.TimeoutExpired:
            return {"status": "error", "reason": "LibreOffice recalc timed out (300s)."}
        cand = os.path.join(td, os.path.splitext(os.path.basename(path))[0] + ".xlsx")
        if not os.path.exists(cand):
            return {"status": "error", "reason": (p.stderr or p.stdout or "conversion produced no file")[:400]}
        wb = load_workbook(cand, data_only=True)
        hits, neg = [], []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and v.strip() in ERRORS:
                        hits.append(f"{ws.title}!{c.coordinate}={v.strip()}")
                    elif isinstance(v, (int, float)) and not isinstance(v, bool) and v < 0:
                        if not has_red(c.number_format):
                            neg.append(f"{ws.title}!{c.coordinate}")
        wb.close()
        return {"status": "errors_found" if hits else "clean",
                "errors": hits[:50], "total_errors": len(hits),
                "negatives_not_red": neg[:50], "total_negatives_not_red": len(neg)}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--no-calc", action="store_true")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()

    rep = {"file": a.path, "static": static_check(a.path)}
    rep["recalc"] = {"status": "skipped", "reason": "--no-calc"} if a.no_calc else recalc_check(a.path)

    fails = list(rep["static"]["fail"])
    if rep["recalc"].get("status") == "errors_found":
        fails.append({"code": "formula_errors",
                      "msg": f"{rep['recalc']['total_errors']} live formula errors: "
                             f"{', '.join(rep['recalc']['errors'][:5])}"})
    if rep["recalc"].get("total_negatives_not_red"):
        fails.append({"code": "negative_not_red_computed",
                      "msg": f"{rep['recalc']['total_negatives_not_red']} formula cells COMPUTE negative but "
                             f"their format has no [Red] section: "
                             f"{', '.join(rep['recalc']['negatives_not_red'][:5])}. Rule 10a."})
    rep["verdict"] = "FAIL" if fails else "PASS"

    if a.json:
        print(json.dumps(rep, indent=2, default=str))
        return 1 if fails else 0

    print(f"\n  {rep['verdict']}  {a.path}")
    print(f"  sheets: {', '.join(rep['static']['sheets'])}")
    for nm, s in rep["static"]["sheets"].items():
        print(f"    {nm:<22} rows={s['rows']:<7} formulas={s['formulas']:<6} constants={s['constants']}")
    print(f"  named ranges: {rep['static']['info']['named_ranges']}   "
          f"fullCalcOnLoad: {rep['static']['info']['fullCalcOnLoad']}")
    print(f"  recalc: {rep['recalc'].get('status')}"
          + (f" — {rep['recalc'].get('reason','')}" if rep["recalc"].get("reason") else ""))
    for f in fails:
        print(f"  FAIL [{f['code']}] {f['msg']}")
    for w in rep["static"]["warn"]:
        print(f"  WARN [{w['code']}] {w['msg']}")
    print()
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
